from pathlib import Path
import docx
import pytest

from app.analyzers.document_analyzer import document_text_extractor
from app.analyzers.document_rules_extractor import document_rules_extractor


def test_txt_document_extraction(tmp_path: Path):
    """Test text extraction and entity recognition from TXT files."""
    sample_text = (
        "Welcome to the Financial Inclusion in Africa Challenge.\n\n"
        "The objective is to predict whether a respondent has a bank_account or not.\n"
        "The unique identifier is 'unique_id'.\n"
        "The evaluation metric for this competition is ROC-AUC (Area Under the ROC Curve).\n"
        "Submissions are evaluated on the Area Under the Curve.\n"
        "Limit of 3 submissions per day."
    )
    doc_path = tmp_path / "competition_overview.txt"
    doc_path.write_text(sample_text, encoding="utf-8")

    chunks = document_text_extractor.extract_text_with_pages(doc_path)
    assert len(chunks) == 1
    assert "Financial Inclusion" in chunks[0][0]

    facts = document_rules_extractor.extract_from_file(doc_path)
    assert "ROC-AUC" in facts.metric_mentions
    assert "bank_account" in facts.target_variable_mentions
    assert "unique_id" in facts.id_column_mentions
    assert len(facts.evidence_snippets) >= 2

    # Verify evidence structure
    metric_evidence = next((e for e in facts.evidence_snippets if "ROC" in e.text_snippet), None)
    assert metric_evidence is not None
    assert metric_evidence.source_file == "competition_overview.txt"
    assert metric_evidence.confidence >= 0.9


def test_docx_document_extraction(tmp_path: Path):
    """Test text extraction and rule extraction from DOCX documents."""
    doc = docx.Document()
    doc.add_heading("Customer Churn Prediction Challenge", 0)
    doc.add_paragraph("The goal is to predict customer 'churn' probability based on telecom telemetry.")
    doc.add_paragraph("Each row is identified by customer_id.")
    doc.add_paragraph("The evaluation metric is Log Loss. Probabilities must be strictly between 0 and 1.")

    docx_path = tmp_path / "rules.docx"
    doc.save(str(docx_path))

    facts = document_rules_extractor.extract_from_file(docx_path)
    assert "Log Loss" in facts.metric_mentions
    assert "churn" in facts.target_variable_mentions
    assert "customer_id" in facts.id_column_mentions


def test_multiple_document_consolidation(tmp_path: Path):
    """Test consolidating facts across multiple documentation files."""
    doc1 = tmp_path / "doc1.txt"
    doc1.write_text("Objective: predict 'sales' volume. Metric is RMSE.", encoding="utf-8")

    doc2 = tmp_path / "doc2.txt"
    doc2.write_text("Unique identifier is 'store_date_id'. Limit of 5 submissions per day.", encoding="utf-8")

    consolidated = document_rules_extractor.consolidate_document_facts([doc1, doc2])
    assert "RMSE" in consolidated.metric_mentions
    assert "sales" in consolidated.target_variable_mentions
    assert "store_date_id" in consolidated.id_column_mentions
    assert len(consolidated.evidence_snippets) >= 2
