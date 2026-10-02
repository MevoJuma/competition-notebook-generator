from pathlib import Path
from typing import List, Tuple
import docx
from pypdf import PdfReader

from app.core.exceptions import DocumentExtractionException
from app.core.logging import get_logger

logger = get_logger("app.analyzers.document_analyzer")


class DocumentTextExtractor:
    """Extracts raw text and page-indexed chunks from PDF, DOCX, and TXT files."""

    @classmethod
    def extract_text_with_pages(cls, file_path: Path) -> List[Tuple[str, int]]:
        """Extract text chunks alongside their 1-indexed page or section numbers."""
        ext = file_path.suffix.lower()

        if not file_path.exists():
            raise DocumentExtractionException(file_path.name, "File does not exist.")

        try:
            if ext == ".pdf":
                return cls._extract_pdf(file_path)
            elif ext in [".docx", ".doc"]:
                return cls._extract_docx(file_path)
            elif ext in [".txt", ".md", ".json"]:
                return cls._extract_text_file(file_path)
            else:
                raise DocumentExtractionException(file_path.name, f"Unsupported document extension '{ext}'.")
        except DocumentExtractionException:
            raise
        except Exception as exc:
            logger.error("Failed to extract text from %s: %s", file_path.name, exc)
            raise DocumentExtractionException(file_path.name, str(exc)) from exc

    @classmethod
    def _extract_pdf(cls, file_path: Path) -> List[Tuple[str, int]]:
        reader = PdfReader(str(file_path))
        chunks: List[Tuple[str, int]] = []
        for idx, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            cleaned = text.strip()
            if cleaned:
                chunks.append((cleaned, idx))
        return chunks

    @classmethod
    def _extract_docx(cls, file_path: Path) -> List[Tuple[str, int]]:
        doc = docx.Document(str(file_path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs)
        return [(full_text, 1)] if full_text else []

    @classmethod
    def _extract_text_file(cls, file_path: Path) -> List[Tuple[str, int]]:
        content = None
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                content = file_path.read_text(encoding=encoding)
                break
            except (UnicodeDecodeError, LookupError):
                continue

        if content is None:
            raise DocumentExtractionException(file_path.name, "Failed to decode text using standard encodings.")

        return [(content.strip(), 1)] if content.strip() else []

    @classmethod
    def extract_full_text(cls, file_path: Path) -> str:
        """Convenience method returning the consolidated text of a document."""
        chunks = cls.extract_text_with_pages(file_path)
        return "\n\n".join(text for text, _ in chunks)


document_text_extractor = DocumentTextExtractor()
